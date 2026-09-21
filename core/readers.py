"""Turn any attachment (txt, xlsx, docx, pdf) into plain text the rule parser and Gemini can read.

Returns a ReadResult: text is None when the file cannot be read as text (e.g. a scanned,
image-only PDF or a corrupt file). `kind` tells the pipeline what to do next:
  text      -> text extracted, parse normally
  scanned   -> image-only; needs Gemini vision (OCR)
  corrupt   -> cannot be opened at all
  unsupported -> unknown format
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass


@dataclass
class ReadResult:
    text: str | None
    kind: str          # text | scanned | corrupt | unsupported
    mime: str
    note: str = ""


MIME = {
    ".txt": "text/plain",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
}


def ext_of(name: str) -> str:
    m = re.search(r"\.[A-Za-z0-9]+$", name or "")
    return m.group(0).lower() if m else ""


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def read_xlsx(data: bytes) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    lines = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [_cell(c) for c in row]
            cells = [c for c in cells if c]
            if not cells:
                continue
            if len(cells) == 1:
                lines.append(cells[0])
            else:
                # key | value [| more]; multi-line party blocks use " | " inside a cell
                key, rest = cells[0], cells[1:]
                val = " ".join(rest)
                parts = [p.strip() for p in val.split(" | ")]
                lines.append(f"{key}: {parts[0]}")
                lines.extend("  " + p for p in parts[1:])
    return "\n".join(lines)


def read_docx(data: bytes) -> str:
    import docx
    d = docx.Document(io.BytesIO(data))
    lines = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for r in t.rows:
            cells = []
            for c in r.cells:  # merged cells repeat; de-duplicate consecutive
                if not cells or cells[-1] != c.text:
                    cells.append(c.text)
            if len(cells) >= 2:
                key, val = cells[0].strip(), "\n".join(cells[1:]).strip()
                vlines = [l.strip() for l in val.splitlines() if l.strip()] or [""]
                lines.append(f"{key}: {vlines[0]}")
                lines.extend("  " + l for l in vlines[1:])
            elif cells:
                lines.append(cells[0])
    return "\n".join(lines)


def _pdf_layout_to_kv(text: str) -> str:
    """pdftotext-style layout: 'Key<2+ spaces>Value' with indented continuation lines."""
    out = []
    for raw in text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            out.append("")
            continue
        if line[:1].isspace():
            out.append("  " + line.strip())
            continue
        m = re.match(r"^(\S.*?)\s{2,}(\S.*)$", line)
        if m and ":" not in m.group(1) and len(m.group(1)) <= 45:
            out.append(f"{m.group(1).strip()}: {m.group(2).strip()}")
        else:
            out.append(line.strip())
    return "\n".join(out)


def _is_bold(w: dict) -> bool:
    return "bold" in (w.get("fontname") or "").lower()


def _page_lines(page) -> list[str]:
    """Rebuild lines from word positions. Bold words at the start of a row are the label
    ('Key: value'); an indented start is a continuation line; wide gaps become 3 spaces.
    Splitting on font also untangles labels that overflow into the value column."""
    words = page.extract_words(keep_blank_chars=False, use_text_flow=False, x_tolerance=1.5,
                               extra_attrs=["fontname"])
    # decorative glyph fonts (check boxes, blots) are not document text
    words = [w for w in words if not re.search(r"dingbat|symbol|wingding", w.get("fontname") or "", re.I)]
    if not words:
        return []
    rows: list[list[dict]] = []
    for w in sorted(words, key=lambda w: (round(w["top"]), w["x0"])):
        if rows and abs(rows[-1][0]["top"] - w["top"]) <= 3:
            rows[-1].append(w)
        else:
            rows.append([w])
    left = min(w["x0"] for w in words)
    lines, prev_bottom = [], None
    for row in rows:
        row.sort(key=lambda w: (not _is_bold(w), w["x0"]))  # bold label words first
        if prev_bottom is not None and row[0]["top"] - prev_bottom > 14:
            lines.append("")
        prev_bottom = max(w["bottom"] for w in row)
        bold = [w for w in row if _is_bold(w)]
        plain = sorted([w for w in row if not _is_bold(w)], key=lambda w: w["x0"])
        if bold and plain and len(bold) <= 8 and min(w["x0"] for w in bold) <= min(w["x0"] for w in plain) + 1:
            key = " ".join(w["text"] for w in sorted(bold, key=lambda w: w["x0"])).rstrip(":")
            lines.append(f"{key}: " + _join(plain))
            continue
        row = sorted(row, key=lambda w: w["x0"])
        char_w = max(1.0, sum((w["x1"] - w["x0"]) / max(1, len(w["text"])) for w in row) / len(row))
        indent = "    " if row[0]["x0"] - left > 6 * char_w else ""
        lines.append(indent + _join(row, char_w))
    return lines


def _join(ws: list[dict], char_w: float | None = None) -> str:
    if char_w is None:
        char_w = max(1.0, sum((w["x1"] - w["x0"]) / max(1, len(w["text"])) for w in ws) / max(1, len(ws)))
    s = ""
    for i, w in enumerate(ws):
        if i:
            s += "   " if w["x0"] - ws[i - 1]["x1"] > 2.5 * char_w else " "
        s += w["text"]
    return s


def read_pdf(data: bytes) -> ReadResult:
    try:
        import pdfplumber
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            lines = [l for p in pdf.pages for l in _page_lines(p) + [""]]
    except Exception as e:  # noqa: BLE001
        return ReadResult(None, "corrupt", MIME[".pdf"], f"PDF could not be opened: {type(e).__name__}")
    joined = "\n".join(lines).strip()
    if len(re.sub(r"\s", "", joined)) < 30:
        return ReadResult(None, "scanned", MIME[".pdf"], "No text layer (image-only / scanned PDF).")
    return ReadResult(_pdf_layout_to_kv(joined), "text", MIME[".pdf"], "Text layer extracted.")


def read_any(name: str, data: bytes) -> ReadResult:
    ext = ext_of(name)
    mime = MIME.get(ext, "application/octet-stream")
    try:
        if ext in ("", ".txt", ".csv", ".md"):
            return ReadResult(data.decode("utf-8", errors="replace"), "text", "text/plain")
        if ext == ".xlsx":
            return ReadResult(read_xlsx(data), "text", mime)
        if ext == ".docx":
            return ReadResult(read_docx(data), "text", mime)
        if ext == ".pdf":
            return read_pdf(data)
        if ext in (".png", ".jpg", ".jpeg"):
            return ReadResult(None, "scanned", mime, "Image file; needs OCR.")
    except Exception as e:  # noqa: BLE001
        return ReadResult(None, "corrupt", mime, f"Could not open {ext} file: {type(e).__name__}: {e}")
    return ReadResult(None, "unsupported", mime, f"Unsupported file type '{ext}'.")
